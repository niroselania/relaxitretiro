import hmac
import io
import os
import secrets
import sqlite3
from datetime import datetime
from functools import wraps

import barcode
from barcode.writer import SVGWriter
from flask import (Flask, abort, flash, g, make_response, redirect, render_template,
                   request, send_file, session, url_for)
from openpyxl import Workbook, load_workbook
from werkzeug.security import check_password_hash, generate_password_hash


app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY", ""),
    DATABASE_PATH=os.environ.get("DATABASE_PATH", "data/deposito.sqlite3"),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("SESSION_COOKIE_SECURE", "false").lower() == "true",
    MAX_CONTENT_LENGTH=12 * 1024 * 1024,
)

SCHEMA = """
CREATE TABLE IF NOT EXISTS users (
  id INTEGER PRIMARY KEY, username TEXT NOT NULL UNIQUE, password_hash TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS containers (
  id INTEGER PRIMARY KEY, code TEXT NOT NULL UNIQUE, name TEXT NOT NULL,
  location TEXT NOT NULL DEFAULT '', site TEXT NOT NULL DEFAULT 'Bariloche', active INTEGER NOT NULL DEFAULT 1,
  parent_id INTEGER REFERENCES containers(id),
  created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS products (
  id INTEGER PRIMARY KEY, sku TEXT NOT NULL UNIQUE, name TEXT NOT NULL,
  item TEXT NOT NULL DEFAULT '', color TEXT NOT NULL DEFAULT '', size TEXT NOT NULL DEFAULT '',
  description TEXT NOT NULL DEFAULT '', legacy_units INTEGER NOT NULL DEFAULT 0,
  source_location TEXT NOT NULL DEFAULT '', active INTEGER NOT NULL DEFAULT 1,
  created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS stock (
  container_id INTEGER NOT NULL REFERENCES containers(id),
  product_id INTEGER NOT NULL REFERENCES products(id), quantity INTEGER NOT NULL CHECK(quantity >= 0),
  PRIMARY KEY(container_id, product_id)
);
CREATE TABLE IF NOT EXISTS movements (
  id INTEGER PRIMARY KEY, created_at TEXT NOT NULL, user_id INTEGER NOT NULL REFERENCES users(id),
  from_container_id INTEGER REFERENCES containers(id), to_container_id INTEGER REFERENCES containers(id),
  from_source TEXT NOT NULL DEFAULT 'external', product_id INTEGER NOT NULL REFERENCES products(id), quantity INTEGER NOT NULL CHECK(quantity > 0),
  note TEXT NOT NULL DEFAULT ''
);
CREATE INDEX IF NOT EXISTS idx_movements_created ON movements(created_at DESC);
CREATE INDEX IF NOT EXISTS idx_stock_product ON stock(product_id);
"""


def now_text():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def get_db():
    if "db" not in g:
        os.makedirs(os.path.dirname(app.config["DATABASE_PATH"]) or ".", exist_ok=True)
        g.db = sqlite3.connect(app.config["DATABASE_PATH"])
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
        g.db.execute("PRAGMA journal_mode = WAL")
    return g.db


@app.teardown_appcontext
def close_db(_error=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = get_db()
    db.executescript(SCHEMA)
    # Keep existing local databases usable when new product and site fields are added.
    for table, column, declaration in [
        ("containers", "site", "TEXT NOT NULL DEFAULT 'Bariloche'"),
        ("containers", "parent_id", "INTEGER REFERENCES containers(id)"),
        ("products", "item", "TEXT NOT NULL DEFAULT ''"),
        ("products", "color", "TEXT NOT NULL DEFAULT ''"),
        ("products", "size", "TEXT NOT NULL DEFAULT ''"),
        ("products", "legacy_units", "INTEGER NOT NULL DEFAULT 0"),
        ("products", "source_location", "TEXT NOT NULL DEFAULT ''"),
        ("products", "active", "INTEGER NOT NULL DEFAULT 1"),
        ("movements", "from_source", "TEXT NOT NULL DEFAULT 'external'"),
    ]:
        columns = {row[1] for row in db.execute(f"PRAGMA table_info({table})")}
        if column not in columns:
            db.execute(f"ALTER TABLE {table} ADD COLUMN {column} {declaration}")
    username = os.environ.get("ADMIN_USERNAME", "admin").strip()
    password = os.environ.get("ADMIN_PASSWORD", "")
    if not app.config["SECRET_KEY"] or not password:
        raise RuntimeError("Configurá SECRET_KEY y ADMIN_PASSWORD antes de iniciar el sistema.")
    if not db.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone():
        db.execute("INSERT INTO users(username, password_hash) VALUES (?, ?)",
                   (username, generate_password_hash(password)))
    db.commit()


with app.app_context():
    init_db()


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapped


@app.before_request
def protect_forms():
    if request.method == "POST":
        sent = request.form.get("csrf_token", "")
        saved = session.get("csrf_token", "")
        if not sent or not saved or not hmac.compare_digest(sent, saved):
            abort(400, "La sesión del formulario venció. Volvé a cargar la pantalla.")


@app.context_processor
def inject_template_values():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(24)
    return {"csrf_token": session["csrf_token"], "active_path": request.path}


@app.template_filter("qty")
def quantity_format(value):
    return f"{value:,}".replace(",", ".")


@app.template_filter("date_short")
def date_short(value):
    try:
        return datetime.fromisoformat(value).strftime("%d/%m/%Y %H:%M")
    except (TypeError, ValueError):
        return value or "—"


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = get_db().execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["csrf_token"] = secrets.token_urlsafe(24)
            return redirect(url_for("dashboard"))
        flash("Usuario o contraseña incorrectos.", "error")
    return render_template("login.html")


@app.post("/logout")
@login_required
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/cuenta", methods=["GET", "POST"])
@login_required
def account():
    if request.method == "POST":
        current = request.form.get("current_password", "")
        new = request.form.get("new_password", "")
        confirm = request.form.get("confirm_password", "")
        user = get_db().execute("SELECT * FROM users WHERE id=?", (session["user_id"],)).fetchone()
        if not check_password_hash(user["password_hash"], current):
            flash("La contraseña actual no es correcta.", "error")
        elif len(new) < 12:
            flash("La nueva contraseña debe tener al menos 12 caracteres.", "error")
        elif new != confirm:
            flash("La confirmación no coincide.", "error")
        else:
            db = get_db()
            db.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(new), session["user_id"]))
            db.commit()
            flash("Contraseña actualizada.", "success")
            return redirect(url_for("account"))
    return render_template("account.html")


@app.route("/")
@login_required
def dashboard():
    db = get_db()
    stats = {
        "containers": db.execute("SELECT COUNT(*) FROM containers WHERE active=1").fetchone()[0],
        "products": db.execute("SELECT COUNT(*) FROM products").fetchone()[0],
        "units": db.execute("SELECT COALESCE(SUM(quantity),0) FROM stock").fetchone()[0],
        "movements": db.execute("SELECT COUNT(*) FROM movements WHERE date(created_at)=date('now','localtime')").fetchone()[0],
    }
    recent = db.execute("""SELECT m.*, p.sku, p.name product,
      CASE WHEN f.name IS NOT NULL THEN f.name WHEN m.from_source='legacy' THEN 'Planilla sin ubicación' WHEN m.from_source='container' THEN 'Retiro manual' ELSE 'Ingreso externo' END from_name,
      t.name to_name
      FROM movements m JOIN products p ON p.id=m.product_id
      LEFT JOIN containers f ON f.id=m.from_container_id
      LEFT JOIN containers t ON t.id=m.to_container_id
      ORDER BY m.id DESC LIMIT 8""").fetchall()
    return render_template("dashboard.html", stats=stats, recent=recent)


def next_container_code(db):
    # Numeric Code 128 identifiers remain stable and are never recycled.
    last = db.execute("SELECT MAX(CAST(code AS INTEGER)) FROM containers WHERE code GLOB '[0-9]*' AND length(code)=12").fetchone()[0]
    candidate = (int(last) + 1) if last else 1
    while db.execute("SELECT 1 FROM containers WHERE code=?", (f"{candidate:012d}",)).fetchone():
        candidate += 1
    return f"{candidate:012d}"


@app.route("/contenedores")
@login_required
def containers():
    q = request.args.get("q", "").strip()
    rows = get_db().execute("""SELECT c.*, COALESCE(SUM(s.quantity),0) units,
      COUNT(DISTINCT s.product_id) products_count
      FROM containers c LEFT JOIN stock s ON s.container_id=c.id
      WHERE c.active=1 AND (?='' OR c.name LIKE ? OR c.code LIKE ? OR c.location LIKE ?)
      GROUP BY c.id ORDER BY c.name COLLATE NOCASE""", (q, f"%{q}%", f"%{q}%", f"%{q}%")).fetchall()
    return render_template("containers.html", containers=rows, q=q)


@app.post("/contenedores/nuevo")
@login_required
def create_container():
    name = request.form.get("name", "").strip()
    location = request.form.get("location", "").strip()
    site = request.form.get("site", "").strip()
    parent_id = request.form.get("parent_id", type=int)
    if not name:
        flash("Ingresá un nombre para el contenedor.", "error")
        return redirect(url_for("containers"))
    if parent_id:
        parent = get_db().execute("SELECT * FROM containers WHERE id=? AND active=1", (parent_id,)).fetchone()
        if not parent:
            flash("El contenedor principal ya no está disponible.", "error")
            return redirect(url_for("containers"))
        site = parent["site"]
    if site not in ("Bariloche", "Retiro"):
        flash("Elegí Bariloche o Retiro como depósito.", "error")
        return redirect(url_for("containers"))
    db = get_db()
    code = next_container_code(db)
    db.execute("INSERT INTO containers(code,name,location,site,parent_id,created_at) VALUES(?,?,?,?,?,?)", (code, name, location, site, parent_id, now_text()))
    db.commit()
    flash(f"Contenedor creado. Código {code}.", "success")
    return redirect(url_for("container_detail", container_id=db.execute("SELECT id FROM containers WHERE code=?", (code,)).fetchone()[0]))


@app.route("/contenedores/<int:container_id>")
@login_required
def container_detail(container_id):
    db = get_db()
    container = db.execute("SELECT * FROM containers WHERE id=? AND active=1", (container_id,)).fetchone()
    if not container:
        abort(404)
    products = db.execute("""SELECT p.*,s.quantity FROM stock s JOIN products p ON p.id=s.product_id
      WHERE s.container_id=? ORDER BY p.name COLLATE NOCASE""", (container_id,)).fetchall()
    children = db.execute("""SELECT c.*,COALESCE(SUM(s.quantity),0) units,COUNT(DISTINCT s.product_id) products_count
      FROM containers c LEFT JOIN stock s ON s.container_id=c.id WHERE c.parent_id=? AND c.active=1
      GROUP BY c.id ORDER BY c.name COLLATE NOCASE""", (container_id,)).fetchall()
    parent = db.execute("SELECT id,name FROM containers WHERE id=?", (container["parent_id"],)).fetchone() if container["parent_id"] else None
    return render_template("container_detail.html", container=container, products=products, children=children, parent=parent)


@app.post("/contenedores/<int:container_id>/productos/<int:product_id>/quitar")
@login_required
def remove_container_product(container_id, product_id):
    db = get_db()
    row = db.execute("SELECT quantity FROM stock WHERE container_id=? AND product_id=?", (container_id, product_id)).fetchone()
    if not row:
        flash("Ese producto ya no está en el contenedor.", "error")
    else:
        db.execute("DELETE FROM stock WHERE container_id=? AND product_id=?", (container_id, product_id))
        db.execute("INSERT INTO movements(created_at,user_id,from_container_id,to_container_id,from_source,product_id,quantity,note) VALUES(?,?,?,NULL,'container',?,?,?)",
                   (now_text(), session["user_id"], container_id, product_id, row["quantity"], "Retiro manual del contenedor"))
        db.commit()
        flash("Producto quitado del contenedor; la salida quedó registrada en movimientos.", "success")
    return redirect(url_for("container_detail", container_id=container_id))


@app.post("/contenedores/<int:container_id>/eliminar")
@login_required
def delete_container(container_id):
    db = get_db()
    container = db.execute("SELECT * FROM containers WHERE id=? AND active=1", (container_id,)).fetchone()
    if not container:
        abort(404)
    subtree = [container_id]
    for cid in subtree:
        subtree.extend(r[0] for r in db.execute("SELECT id FROM containers WHERE parent_id=? AND active=1", (cid,)).fetchall())
    placeholders = ",".join("?" for _ in subtree)
    if request.form.get("confirm") != "yes":
        abort(400)
    # Deleting a hierarchy also removes its remaining stock, recorded as outgoing adjustments.
    for cid in subtree:
        stock_rows = db.execute("SELECT product_id,quantity FROM stock WHERE container_id=?", (cid,)).fetchall()
        for row in stock_rows:
            db.execute("INSERT INTO movements(created_at,user_id,from_container_id,to_container_id,from_source,product_id,quantity,note) VALUES(?,?,?,NULL,'container',?,?,?)",
                       (now_text(), session["user_id"], cid, row["product_id"], row["quantity"], "Eliminación de contenedor"))
    db.execute(f"DELETE FROM stock WHERE container_id IN ({placeholders})", subtree)
    db.execute(f"UPDATE containers SET active=0 WHERE id IN ({placeholders})", subtree)
    db.commit()
    flash("Contenedor y sus subcontenedores eliminados; las cantidades se registraron como salidas.", "success")
    return redirect(url_for("containers"))


@app.route("/contenedores/<int:container_id>/etiqueta.svg")
@login_required
def container_label(container_id):
    container = get_db().execute("SELECT * FROM containers WHERE id=? AND active=1", (container_id,)).fetchone()
    if not container:
        abort(404)
    code128 = barcode.get("code128", container["code"], writer=SVGWriter())
    svg = code128.render({"module_width": 0.28, "module_height": 14, "quiet_zone": 3,
                          "font_size": 10, "text_distance": 2, "write_text": True})
    response = make_response(svg)
    response.headers["Content-Type"] = "image/svg+xml; charset=utf-8"
    response.headers["Content-Disposition"] = f'inline; filename="contenedor-{container["code"]}.svg"'
    return response


@app.route("/contenedores/<int:container_id>/imprimir")
@login_required
def print_container_label(container_id):
    container = get_db().execute("SELECT * FROM containers WHERE id=?", (container_id,)).fetchone()
    if not container:
        abort(404)
    return render_template("label.html", container=container)


@app.route("/productos")
@login_required
def products():
    code = request.args.get("codigo", "").strip()
    q = request.args.get("q", "").strip()
    requested_sites = request.args.getlist("deposito")
    sites = [site for site in ("Bariloche", "Retiro") if site in requested_sites]
    if not request.args:
        sites = ["Bariloche", "Retiro"]
    db = get_db()
    archived = request.args.get("archivados") == "1"
    if archived:
        rows = db.execute("""SELECT p.*,p.legacy_units+COALESCE(SUM(s.quantity),0) units,
          COALESCE(GROUP_CONCAT(c.name || ' (' || s.quantity || ')', ' · '),p.source_location,'') locations
          FROM products p LEFT JOIN stock s ON s.product_id=p.id AND s.quantity>0
          LEFT JOIN containers c ON c.id=s.container_id
          WHERE p.active=0 AND (?='' OR p.sku LIKE ? OR p.name LIKE ? OR p.item LIKE ? OR p.color LIKE ? OR p.size LIKE ? OR p.description LIKE ?)
            AND (?='' OR p.sku LIKE ?)
          GROUP BY p.id ORDER BY p.sku COLLATE NOCASE""",
          (q,f"%{q}%",f"%{q}%",f"%{q}%",f"%{q}%",f"%{q}%",f"%{q}%",code,f"%{code}%")).fetchall()
    elif len(sites) == 2:
        placeholders = ",".join("?" for _ in sites)
        rows = db.execute(f"""SELECT p.*,p.legacy_units+COALESCE(SUM(CASE WHEN c.id IS NOT NULL THEN s.quantity ELSE 0 END),0) units,
          CASE WHEN p.legacy_units>0 THEN 'Pendiente de ubicación (' || p.legacy_units || ')'
            || CASE WHEN GROUP_CONCAT(CASE WHEN c.id IS NOT NULL THEN c.name || ' (' || s.quantity || ')' END, ' · ') IS NOT NULL
              THEN ' · ' || GROUP_CONCAT(CASE WHEN c.id IS NOT NULL THEN c.name || ' (' || s.quantity || ')' END, ' · ') ELSE '' END
            ELSE COALESCE(GROUP_CONCAT(CASE WHEN c.id IS NOT NULL THEN c.name || ' (' || s.quantity || ')' END, ' · '),NULLIF(p.source_location,''),'') END locations
          FROM products p LEFT JOIN stock s ON s.product_id=p.id AND s.quantity>0
          LEFT JOIN containers c ON c.id=s.container_id AND c.active=1 AND c.site IN ({placeholders})
          WHERE p.active=1
            AND (?='' OR p.sku LIKE ? OR p.name LIKE ? OR p.item LIKE ? OR p.color LIKE ? OR p.size LIKE ? OR p.description LIKE ? OR c.name LIKE ? OR c.location LIKE ?)
            AND (?='' OR p.sku LIKE ?)
          GROUP BY p.id ORDER BY p.sku COLLATE NOCASE""",
          (*sites, q, f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%", code, f"%{code}%")).fetchall()
    elif len(sites) == 1:
        rows = db.execute("""SELECT p.*,COALESCE(SUM(s.quantity),0) units,
          GROUP_CONCAT(c.name || ' (' || s.quantity || ')', ' · ') locations
          FROM products p JOIN stock s ON s.product_id=p.id AND s.quantity>0
          JOIN containers c ON c.id=s.container_id AND c.active=1 AND c.site=?
          WHERE p.active=1
            AND (?='' OR p.sku LIKE ? OR p.name LIKE ? OR p.item LIKE ? OR p.color LIKE ? OR p.size LIKE ? OR p.description LIKE ? OR c.name LIKE ? OR c.location LIKE ?)
            AND (?='' OR p.sku LIKE ?)
          GROUP BY p.id ORDER BY p.sku COLLATE NOCASE""",
          (sites[0],q,f"%{q}%",f"%{q}%",f"%{q}%",f"%{q}%",f"%{q}%",f"%{q}%",f"%{q}%",f"%{q}%",code,f"%{code}%")).fetchall()
    else:
        rows = []
    return render_template("products.html", products=rows, q=q, code=code, sites=sites, archived=archived)


@app.post("/productos/<int:product_id>/archivar")
@login_required
def archive_product(product_id):
    db = get_db()
    product = db.execute("SELECT id FROM products WHERE id=? AND active=1", (product_id,)).fetchone()
    if not product:
        abort(404)
    held = db.execute("SELECT COALESCE(SUM(quantity),0) FROM stock WHERE product_id=?", (product_id,)).fetchone()[0]
    if held:
        flash("No se puede archivar: todavía hay unidades ubicadas en contenedores.", "error")
    else:
        db.execute("UPDATE products SET active=0,legacy_units=0 WHERE id=?", (product_id,))
        db.commit()
        flash("Producto archivado. Podés recuperarlo desde Productos archivados.", "success")
    return redirect(url_for("products", deposito=["Retiro", "Bariloche"]))


@app.post("/productos/<int:product_id>/restaurar")
@login_required
def restore_product(product_id):
    db = get_db()
    if not db.execute("SELECT 1 FROM products WHERE id=? AND active=0", (product_id,)).fetchone():
        abort(404)
    db.execute("UPDATE products SET active=1 WHERE id=?", (product_id,))
    db.commit()
    flash("Producto recuperado del archivo.", "success")
    return redirect(url_for("products", deposito=["Retiro", "Bariloche"]))


@app.post("/productos/nuevo")
@login_required
def create_product():
    sku = request.form.get("sku", "").strip()
    name = request.form.get("name", "").strip()
    item = request.form.get("item", "").strip()
    color = request.form.get("color", "").strip()
    size = request.form.get("size", "").strip()
    description = request.form.get("description", "").strip()
    if not sku or not name:
        flash("El código/SKU y el nombre son obligatorios.", "error")
        return redirect(url_for("products"))
    try:
        db = get_db()
        db.execute("INSERT INTO products(sku,name,item,color,size,description,created_at) VALUES(?,?,?,?,?,?,?)", (sku, name, item, color, size, description, now_text()))
        db.commit()
    except sqlite3.IntegrityError:
        flash("Ese código de producto ya existe.", "error")
    else:
        flash("Producto agregado.", "success")
    return redirect(url_for("products"))


@app.route("/stock")
@login_required
def stock_view():
    q = request.args.get("q", "").strip()
    rows = get_db().execute("""SELECT c.id container_id,c.code container_code,c.name container_name,c.location,
      p.sku,p.name product_name,s.quantity
      FROM stock s JOIN containers c ON c.id=s.container_id JOIN products p ON p.id=s.product_id
      WHERE c.active=1 AND s.quantity>0 AND (?='' OR p.name LIKE ? OR p.sku LIKE ? OR c.name LIKE ? OR c.code LIKE ?)
      ORDER BY c.name COLLATE NOCASE,p.name COLLATE NOCASE""", (q,f"%{q}%",f"%{q}%",f"%{q}%",f"%{q}%")).fetchall()
    return render_template("stock.html", rows=rows, q=q)


@app.route("/lector")
@login_required
def scan():
    code = request.args.get("codigo", "").strip()
    db = get_db()
    container = db.execute("SELECT * FROM containers WHERE code=? AND active=1", (code,)).fetchone() if code else None
    product = db.execute("SELECT * FROM products WHERE sku=? AND active=1", (code,)).fetchone() if code else None
    container_stock = []
    product_stock = []
    if container:
        container_stock = db.execute("""SELECT p.sku,p.name,s.quantity FROM stock s JOIN products p ON p.id=s.product_id
          WHERE s.container_id=? AND s.quantity>0 ORDER BY p.name COLLATE NOCASE""", (container["id"],)).fetchall()
    if product:
        product_stock = db.execute("""SELECT c.id,c.name,c.code,c.location,s.quantity FROM stock s JOIN containers c ON c.id=s.container_id
          WHERE s.product_id=? AND s.quantity>0 ORDER BY c.name COLLATE NOCASE""", (product["id"],)).fetchall()
    return render_template("scan.html", code=code, container=container, product=product,
                           container_stock=container_stock, product_stock=product_stock)


@app.route("/movimientos", methods=["GET", "POST"])
@login_required
def movements():
    db = get_db()
    if request.method == "POST":
        product_id = request.form.get("product_id", type=int)
        from_value = request.form.get("from_id", "")
        legacy_source = from_value == "legacy"
        from_id = None if legacy_source else (request.form.get("from_id", type=int) or None)
        to_id = request.form.get("to_id", type=int) or None
        quantity = request.form.get("quantity", type=int)
        note = request.form.get("note", "").strip()
        product = db.execute("SELECT id,legacy_units FROM products WHERE id=? AND active=1", (product_id,)).fetchone()
        if not product or not quantity or quantity < 1 or (from_id is None and not legacy_source and to_id is None):
            flash("Revisá producto, cantidad, origen y destino.", "error")
        elif legacy_source and (not to_id or product["legacy_units"] < quantity):
            flash("La cantidad supera las unidades importadas pendientes o falta elegir el contenedor de destino.", "error")
        elif from_id and from_id == to_id:
            flash("El origen y el destino deben ser distintos.", "error")
        elif from_id and not db.execute("SELECT id FROM containers WHERE id=? AND active=1", (from_id,)).fetchone():
            flash("El contenedor de origen no existe o está inactivo.", "error")
        elif to_id and not db.execute("SELECT id FROM containers WHERE id=? AND active=1", (to_id,)).fetchone():
            flash("El contenedor de destino no existe o está inactivo.", "error")
        else:
            try:
                with db:
                    if legacy_source:
                        db.execute("UPDATE products SET legacy_units=legacy_units-? WHERE id=?", (quantity,product_id))
                    elif from_id:
                        row = db.execute("SELECT quantity FROM stock WHERE container_id=? AND product_id=?", (from_id, product_id)).fetchone()
                        if not row or row["quantity"] < quantity:
                            raise ValueError("No hay esa cantidad disponible en el origen.")
                        remaining = row["quantity"] - quantity
                        if remaining:
                            db.execute("UPDATE stock SET quantity=? WHERE container_id=? AND product_id=?", (remaining, from_id, product_id))
                        else:
                            db.execute("DELETE FROM stock WHERE container_id=? AND product_id=?", (from_id, product_id))
                    if to_id:
                        db.execute("""INSERT INTO stock(container_id,product_id,quantity) VALUES(?,?,?)
                          ON CONFLICT(container_id,product_id) DO UPDATE SET quantity=quantity+excluded.quantity""", (to_id, product_id, quantity))
                    db.execute("""INSERT INTO movements(created_at,user_id,from_container_id,to_container_id,from_source,product_id,quantity,note)
                      VALUES(?,?,?,?,?,?,?,?)""", (now_text(),session["user_id"],from_id,to_id,"legacy" if legacy_source else "external",product_id,quantity,note))
                flash("Movimiento registrado.", "success")
                return redirect(url_for("movements"))
            except ValueError as exc:
                flash(str(exc), "error")
    rows = db.execute("""SELECT m.*,p.sku,p.name product,
      CASE WHEN f.name IS NOT NULL THEN f.name WHEN m.from_source='legacy' THEN 'Planilla sin ubicación' WHEN m.from_source='container' THEN 'Retiro manual' ELSE 'Ingreso externo' END from_name,
      t.name to_name
      FROM movements m JOIN products p ON p.id=m.product_id
      LEFT JOIN containers f ON f.id=m.from_container_id LEFT JOIN containers t ON t.id=m.to_container_id
      ORDER BY m.id DESC LIMIT 100""").fetchall()
    return render_template("movements.html", rows=rows,
      products=db.execute("SELECT * FROM products WHERE active=1 ORDER BY name COLLATE NOCASE").fetchall(),
      containers=db.execute("SELECT * FROM containers WHERE active=1 ORDER BY name COLLATE NOCASE").fetchall())


def import_elementos_sheet(db, sheet):
    headers = [str(cell or "").strip().lower() for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    required = ["codigo", "nombre", "item", "color", "talle", "descripcion", "unidades", "ubicaciones"]
    if not set(required).issubset(headers):
        raise ValueError("La hoja elementos debe incluir codigo, nombre, item, color, talle, descripcion, unidades y ubicaciones.")
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if any(v is not None for v in row)]
    seen = set()
    total_units = 0
    cleaned = []
    for row_number, row in enumerate(rows, start=2):
        code = str(row.get("codigo") or "").strip()
        name = str(row.get("nombre") or "").strip()
        if not code or not name:
            raise ValueError(f"Falta código o nombre en la fila {row_number}.")
        if code in seen:
            raise ValueError(f"El código {code} está repetido en la planilla.")
        seen.add(code)
        raw_units = row.get("unidades")
        try:
            units = int(raw_units or 0)
            if units < 0 or (raw_units is not None and float(raw_units) != units):
                raise ValueError
        except (TypeError, ValueError):
            raise ValueError(f"La cantidad de unidades no es un entero válido en la fila {row_number}.")
        total_units += units
        cleaned.append((code,name,str(row.get("item") or "").strip(),str(row.get("color") or "").strip(),
          str(row.get("talle") or "").strip(),str(row.get("descripcion") or "").strip(),units,
          str(row.get("ubicaciones") or "").strip()))
    with db:
        for item in cleaned:
            db.execute("""INSERT INTO products(sku,name,item,color,size,description,legacy_units,source_location,created_at)
              VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(sku) DO UPDATE SET
              name=excluded.name,item=excluded.item,color=excluded.color,size=excluded.size,
              description=excluded.description,source_location=excluded.source_location,active=1""",
              (*item,now_text()))
    location_count = sum(bool(item[7]) for item in cleaned)
    return len(cleaned), total_units, location_count


@app.route("/importar", methods=["GET", "POST"])
@login_required
def import_data():
    if request.method == "POST":
        upload = request.files.get("file")
        if not upload or not upload.filename.lower().endswith((".xlsx", ".xlsm")):
            flash("Subí un archivo Excel .xlsx con las hojas Contenedores y Stock.", "error")
            return redirect(url_for("import_data"))
        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
            if "elementos" in workbook.sheetnames:
                db = get_db()
                count, total_units, location_count = import_elementos_sheet(db, workbook["elementos"])
                flash(f"Importación lista: {count:,} productos y {total_units:,} unidades. Las ubicaciones informadas cubren {location_count} filas; lo demás queda pendiente de ubicar.".replace(",", "."), "success")
                return redirect(url_for("import_data"))
            if not {"Contenedores", "Productos", "Stock"}.issubset(workbook.sheetnames):
                raise ValueError("No encontré la hoja elementos ni la plantilla con hojas Contenedores, Productos y Stock.")
            db = get_db()
            parsed = {}
            for sheet_name, required in {
                "Contenedores": ["codigo", "nombre", "ubicacion", "sucursal"],
                "Productos": ["sku", "nombre", "item", "color", "talle", "descripcion"],
                "Stock": ["codigo_contenedor", "sku", "cantidad"],
            }.items():
                sheet = workbook[sheet_name]
                headers = [str(c or "").strip().lower() for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
                if not set(required).issubset(headers):
                    raise ValueError(f"La hoja {sheet_name} debe incluir columnas: {', '.join(required)}.")
                parsed[sheet_name] = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if any(v is not None for v in row)]
            container_codes = {str(r["codigo"]).strip() for r in parsed["Contenedores"] if r.get("codigo") and r.get("nombre")}
            product_skus = {str(r["sku"]).strip() for r in parsed["Productos"] if r.get("sku") and r.get("nombre")}
            for r in parsed["Stock"]:
                if str(r.get("codigo_contenedor", "")).strip() not in container_codes or str(r.get("sku", "")).strip() not in product_skus:
                    raise ValueError("Cada fila de Stock debe apuntar a un código de contenedor y SKU incluidos en el mismo Excel.")
                if not isinstance(r.get("cantidad"), (int, float)) or int(r["cantidad"]) != r["cantidad"] or r["cantidad"] < 0:
                    raise ValueError("Las cantidades de Stock deben ser números enteros iguales o mayores que cero.")
            if len(container_codes) != len(parsed["Contenedores"]) or len(product_skus) != len(parsed["Productos"]):
                raise ValueError("Hay códigos o SKU vacíos o repetidos en el archivo.")
            with db:
                for r in parsed["Contenedores"]:
                    site = str(r.get("sucursal") or r.get("deposito") or "").strip().title()
                    if site not in ("Bariloche", "Retiro"):
                        raise ValueError("La sucursal de cada contenedor debe ser Bariloche o Retiro.")
                    db.execute("""INSERT INTO containers(code,name,location,site,created_at) VALUES(?,?,?,?,?)
                      ON CONFLICT(code) DO UPDATE SET name=excluded.name,location=excluded.location,site=excluded.site,active=1""",
                      (str(r["codigo"]).strip(),str(r["nombre"]).strip(),str(r.get("ubicacion") or "").strip(),site,now_text()))
                for r in parsed["Productos"]:
                    pending = r.get("unidades_sin_ubicar") or 0
                    if not isinstance(pending, (int,float)) or int(pending) != pending or pending < 0:
                        raise ValueError("Las unidades sin ubicar deben ser números enteros iguales o mayores que cero.")
                    db.execute("""INSERT INTO products(sku,name,item,color,size,description,legacy_units,source_location,created_at)
                      VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(sku) DO UPDATE SET
                      name=excluded.name,item=excluded.item,color=excluded.color,size=excluded.size,
                      description=excluded.description,source_location=excluded.source_location,active=1""",
                      (str(r["sku"]).strip(),str(r["nombre"]).strip(),str(r.get("item") or "").strip(),str(r.get("color") or "").strip(),str(r.get("talle") or r.get("size") or "").strip(),str(r.get("descripcion") or "").strip(),int(pending),str(r.get("ubicaciones") or "").strip(),now_text()))
                for r in parsed["Stock"]:
                    cid = db.execute("SELECT id FROM containers WHERE code=?", (str(r["codigo_contenedor"]).strip(),)).fetchone()[0]
                    pid = db.execute("SELECT id FROM products WHERE sku=?", (str(r["sku"]).strip(),)).fetchone()[0]
                    db.execute("""INSERT INTO stock(container_id,product_id,quantity) VALUES(?,?,?)
                      ON CONFLICT(container_id,product_id) DO UPDATE SET quantity=excluded.quantity""", (cid,pid,int(r["cantidad"])))
            flash(f"Importación lista: {len(parsed['Contenedores'])} contenedores, {len(parsed['Productos'])} productos, {len(parsed['Stock'])} existencias. Los códigos importados se conservaron.", "success")
            return redirect(url_for("import_data"))
        except (ValueError, KeyError, TypeError, sqlite3.IntegrityError) as exc:
            flash(str(exc), "error")
        except Exception:
            app.logger.exception("Error al procesar archivo de importación")
            flash("No se pudo leer el archivo. Descargá la plantilla y comprobá que sea un .xlsx válido.", "error")
    return render_template("import.html")


@app.get("/importar/plantilla.xlsx")
@login_required
def import_template():
    wb = Workbook()
    sheets = {
        "Contenedores": ["codigo", "nombre", "ubicacion", "sucursal"],
        "Productos": ["sku", "nombre", "item", "color", "talle", "descripcion", "unidades_sin_ubicar", "ubicaciones"],
        "Stock": ["codigo_contenedor", "sku", "cantidad"],
    }
    for index, (name, headers) in enumerate(sheets.items()):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = name
        ws.append(headers)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="plantilla_deposito.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/exportar.xlsx")
@login_required
def export_data():
    db = get_db()
    wb = Workbook()
    for index, (name, query) in enumerate([
        ("Contenedores", "SELECT code codigo,name nombre,location ubicacion,site sucursal FROM containers WHERE active=1 ORDER BY code"),
        ("Productos", "SELECT sku,name nombre,item,color,size talle,description descripcion,legacy_units unidades_sin_ubicar,source_location ubicaciones FROM products ORDER BY sku"),
        ("Stock", "SELECT c.code codigo_contenedor,p.sku,s.quantity cantidad FROM stock s JOIN containers c ON c.id=s.container_id JOIN products p ON p.id=s.product_id WHERE s.quantity>0 ORDER BY c.code,p.sku"),
    ]):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = name
        rows = db.execute(query).fetchall()
        ws.append(list(rows[0].keys()) if rows else {"Contenedores": ["codigo","nombre","ubicacion","sucursal"],"Productos":["sku","nombre","item","color","talle","descripcion","unidades_sin_ubicar","ubicaciones"],"Stock":["codigo_contenedor","sku","cantidad"]}[name])
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="deposito_stock.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.errorhandler(413)
def too_large(_error):
    return "El Excel supera el límite de 12 MB.", 413


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "8000")))
